"""Official fallow figures: MAPA yearbook by province and Aragon CAP declarations by municipality.

Both are parsed into small CSVs under data/official so the rest of the pipeline never touches
the spreadsheets again.
"""
from __future__ import annotations

import io
import re
import unicodedata

import openpyxl
import pandas as pd
import requests

from . import config, geo

# MAPA statistical yearbook, table 3.4 (land under crops by main use). Recent editions publish
# each table as its own file; older ones bundle the whole chapter 3 in one workbook.
ANUARIO_URLS = {
    2025: "https://www.mapa.gob.es/estadistica/pags/anuario/2025/FORMATO%20EXCEL/AE25-C03/3.4.xlsx",
    2024: "https://www.mapa.gob.es/estadistica/pags/anuario/2024/FORMATO%20EXCEL/AE24-C03/3.4.xlsx",
    2023: "https://www.mapa.gob.es/estadistica/pags/anuario/2023-Avance/FORMATO%20EXCEL/AE23-C03/3.4.xlsx",
    2022: "https://www.mapa.gob.es/estadistica/pags/anuario/2022/FORMATO%20EXCEL/AE22-C03/3.4.xlsx",
    2021: "https://www.mapa.gob.es/estadistica/pags/anuario/2021/CAPITULOS%20XLS/AE21-C03.xlsx",
}
ANUARIO_ROW = "barbechos y otras tierras no ocupadas"

# Aragon: surface declared in the CAP single application, by municipality and crop, 2018-2022.
ARAGON_PAC_URL = ("https://www.aragon.es/documents/20127/45811930/"
                  "PAC+2022+Superficie+declarada+por+cultivos.xlsx/ac04eea9-aa37-2b43-19a6-86ff0bd74a69")
# Fallow products in the declarations. Fallow for honey plants and environmental fallow are a
# few hundred hectares and are kept for completeness.
ARAGON_FALLOW = re.compile(r"^BARBECHO", re.I)


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"[^a-z ]", " ", s)
    return " ".join(s.split())


# Yearbook province labels are "Coruña, A", "Rioja, La", "Balears (Illes)"... Match by the
# distinctive word rather than the full label.
_PROVINCE_KEYS = {code: _norm(name).replace("araba ", "").split()[-1]
                  for code, name in geo.PROVINCE_NAMES.items()}
_PROVINCE_KEYS.update({"01": "alava", "07": "balears", "15": "coruna", "20": "gipuzkoa",
                       "26": "rioja", "48": "bizkaia", "29": "malaga", "13": "ciudad real"})
_ALIASES = {"guipuzcoa": "20", "vizcaya": "48", "araba": "01", "asturias": "33",
            "cantabria": "39", "navarra": "31", "madrid": "28", "murcia": "30", "baleares": "07"}
UNIPROVINCIAL = {"07", "26", "28", "30", "31", "33", "39"}


def province_code(label: str) -> str | None:
    n = _norm(label)
    for alias, code in _ALIASES.items():
        if alias in n.split():
            return code
    hits = [c for c, k in _PROVINCE_KEYS.items() if k in n.split() or (" " in k and k in n)]
    return hits[0] if len(hits) == 1 else None


def _download(url: str, name: str) -> bytes:
    path = config.CACHE / name
    if not path.exists():
        config.CACHE.mkdir(parents=True, exist_ok=True)
        r = requests.get(url, timeout=300, headers={"User-Agent": "barbechos-esteparias"})
        r.raise_for_status()
        path.write_bytes(r.content)
    return path.read_bytes()


# --- Yearbook ---------------------------------------------------------------------------------

def parse_anuario(content: bytes) -> pd.DataFrame:
    """Rows (year, province, secano_ha, regadio_ha, total_ha) of the fallow columns of table 3.4."""
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = "3.4" if "3.4" in wb.sheetnames else wb.sheetnames[0]
    rows = list(wb[sheet].iter_rows(values_only=True))
    year = header = col = None
    for i, r in enumerate(rows):
        text = [str(c) for c in r if c is not None]
        if year is None:
            years = [re.findall(r"\b(?:19|20)\d{2}\b", t) for t in text if t.startswith("3.4")]
            if years and years[0]:
                year = int(years[0][-1])
        if col is None:
            for j, c in enumerate(r):
                if c is not None and ANUARIO_ROW in _norm(c):
                    header, col = i, j
                    break
    if year is None or col is None:
        raise ValueError("table 3.4 layout not recognised")
    out = []
    for r in rows[header + 1:]:
        label = r[0]
        if not isinstance(label, str) or not isinstance(r[col], (int, float)):
            continue
        # Communities are in capitals (and, in recent editions, indented one step more).
        # Single-province communities only have that capitalised row, which is the province.
        code = province_code(label)
        if label.strip().isupper() and code not in UNIPROVINCIAL:
            code = None
        if code is None:
            if _norm(label) == "espana":
                code = "ES"
            else:
                continue
        out.append({"year": year, "province": code, "secano_ha": r[col], "regadio_ha": r[col + 1],
                    "total_ha": r[col + 2]})
    return pd.DataFrame(out).drop_duplicates(["year", "province"])


def anuario() -> pd.DataFrame:
    frames = []
    for edition, url in sorted(ANUARIO_URLS.items()):
        frames.append(parse_anuario(_download(url, f"anuario-{edition}-3.4.xlsx")))
    return pd.concat(frames).drop_duplicates(["year", "province"], keep="last").sort_values(["year", "province"])


# --- Aragon CAP declarations -------------------------------------------------------------------

def parse_aragon_pac(content: bytes) -> pd.DataFrame:
    """Rows (year, ine, rainfed_ha, irrigated_ha, declared_ha) of declared fallow per municipality.

    ``declared_ha`` is all the declared surface of the municipality (any product), the
    denominator for fallow shares. From 2022 on, only main crops count, as in the earlier years.
    """
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out = []
    for sheet in wb.sheetnames:
        m = re.search(r"(\d{4})", sheet)
        if not m:
            continue
        rows = wb[sheet].iter_rows(values_only=True)
        header = None
        for r in rows:
            if r and r[0] == "Provincia":
                header = [_norm(c) if c else "" for c in r]
                break
        if header is None:
            continue
        ix = {k: header.index(k) for k in header if k}
        i_ine, i_prod, i_irr, i_ha = ix["municipio ine"], ix["descripcion"], ix["secano regadio"], ix["superficie"]
        i_main = ix.get("principal secundario")
        recs = []
        for r in rows:
            if r[i_ine] is None or not isinstance(r[i_ha], (int, float)):
                continue
            if i_main is not None and str(r[i_main]).strip().lower() != "principal":
                continue
            recs.append((f"{int(r[i_ine]):05d}", str(r[i_prod]), str(r[i_irr]).lower().startswith("sec"), r[i_ha]))
        df = pd.DataFrame(recs, columns=["ine", "product", "rainfed", "ha"])
        fallow = df[df["product"].str.match(ARAGON_FALLOW)]
        g = pd.DataFrame({
            "rainfed_ha": fallow[fallow.rainfed].groupby("ine").ha.sum(),
            "irrigated_ha": fallow[~fallow.rainfed].groupby("ine").ha.sum(),
            "declared_ha": df.groupby("ine").ha.sum(),
        }).fillna(0.0)
        g.insert(0, "year", int(m.group(1)))
        out.append(g.reset_index(names="ine"))
    return pd.concat(out, ignore_index=True)


def aragon_pac() -> pd.DataFrame:
    return parse_aragon_pac(_download(ARAGON_PAC_URL, "aragon-pac-2018-2022.xlsx"))


def update() -> dict[str, int]:
    """Download and parse every official source into data/official."""
    config.OFFICIAL_DIR.mkdir(parents=True, exist_ok=True)
    a = anuario().round(0)
    a.to_csv(config.OFFICIAL_DIR / "anuario-provincias.csv", index=False)
    p = aragon_pac().round(2)
    p.to_csv(config.OFFICIAL_DIR / "aragon-pac-municipios.csv", index=False)
    return {"anuario_rows": len(a), "aragon_rows": len(p)}
